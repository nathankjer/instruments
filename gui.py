"""
Panasonic DC/DC Test Dashboard
------------------------------
A Tkinter-based control panel for the DS1000Z oscilloscope.
Features:
- Real Hardware Control & Simulation Mode
- Safety Interlocks (Voltage Limits, Emergency Stop)
- Master Excel Logging (auto-appends to Panasonic_Master_Log.xlsx)
"""

import tkinter as tk
from tkinter import messagebox, scrolledtext
import json
import datetime
import time
import os
import openpyxl 
from openpyxl import Workbook, load_workbook

# Import the REAL driver
from ds1000z import DS1000Z

# --- MOCK DRIVER (For Simulation) ---
class MockDS1000Z:
    """A fake scope that logs commands to console for offline testing."""
    def __init__(self, ip_address):
        time.sleep(0.5) 
        self.ip = ip_address
        print(f"[SIM] Connected to virtual scope at {ip_address}")
        
    def reset(self):
        print(f"[SIM] {self.ip} > *RST")
        
    def set_source_amplitude(self, volts):
        print(f"[SIM] {self.ip} > :SOUR:VOLT {volts}")
        
    def set_source_frequency(self, freq):
        print(f"[SIM] {self.ip} > :SOUR:FREQ {freq}")
        
    def enable_source(self):
        print(f"[SIM] {self.ip} > :OUTP ON")

# Global variables
scope = None
EXCEL_FILENAME = "Panasonic_Master_Log.xlsx"

def log_message(message):
    """Appends a timestamped entry to the scrolling log console."""
    timestamp = datetime.datetime.now().strftime("%H:%M:%S")
    log_window.configure(state='normal')
    log_window.insert(tk.END, f"[{timestamp}] {message}\n")
    log_window.see(tk.END)
    log_window.configure(state='disabled')

def load_config():
    """Initializes the UI with network settings from config.json."""
    try:
        with open('config.json', 'r') as f:
            config = json.load(f)
        ip_entry.insert(0, config.get('oscilloscope_ip', ''))
        log_message("System Initialized. Configuration loaded.")
    except FileNotFoundError:
        log_message("System Alert: config.json not found.")

def connect_to_scope():
    """Establishes connection (Real or Simulated) and unlocks controls."""
    global scope
    target_ip = ip_entry.get()
    is_sim = sim_mode_var.get()
    
    mode_text = "SIMULATION" if is_sim else "HARDWARE"
    log_message(f"Initiating {mode_text} handshake with {target_ip}...")
    
    try:
        if is_sim:
            scope = MockDS1000Z(target_ip)
            status_indicator.config(bg="#2196F3")
            status_label.config(text="SIMULATED", fg="#2196F3")
        else:
            scope = DS1000Z(target_ip)
            scope.reset()
            status_indicator.config(bg="#4CAF50")
            status_label.config(text="CONNECTED", fg="#4CAF50")
        
        # Unlock Interface
        run_btn.config(state="normal")
        stop_btn.config(state="normal")
        save_btn.config(state="normal")
        log_message(f"Connection established ({mode_text}). Ready.")
        
    except Exception as e:
        status_indicator.config(bg="#F44336")
        status_label.config(text="DISCONNECTED", fg="#F44336")
        log_message(f"Handshake Failed: {e}")
        messagebox.showerror("Connection Error", str(e))

def run_test_sequence():
    """Validates inputs and executes the defined waveform parameters."""
    if not scope:
        messagebox.showerror("Hardware Error", "No instrument connected.")
        return

    # Compliance Check
    tech_name = tech_entry.get().strip()
    if not tech_name:
        messagebox.showwarning("Compliance Error", "Technician Name is required.")
        return

    # Parameter Validation
    try:
        volts = float(volts_entry.get())
        freq = float(freq_entry.get())
    except ValueError:
        messagebox.showerror("Input Error", "Voltage and Frequency must be numeric values.")
        return

    # Safety Interlock
    if volts > 20:
        log_message("SAFETY INTERLOCK: Voltage > 20V rejected.")
        return

    log_message(f"--- SESSION STARTED: {tech_name.upper()} ---")
    log_message(f"Setting Source: {volts}V @ {freq}Hz")
    
    try:
        scope.set_source_amplitude(volts)
        scope.set_source_frequency(freq)
        scope.enable_source()
        log_message("Output Active. Test sequence running.")
        log_message("--- SEQUENCE COMPLETE ---")
    except Exception as e:
        log_message(f"Command Error: {e}")

def disable_output():
    """Emergency Stop: Resets instrument state immediately."""
    if not scope: return
    try:
        scope.reset()
        log_message("SAFETY STOP: Output disabled / Instrument Reset.")
    except Exception as e:
        log_message(f"Stop Failed: {e}")

def update_excel_log():
    """Appends current session data to the master Excel file."""
    
    # 1. Gather Data Points
    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    tech = tech_entry.get().strip() or "Unknown"
    ip = ip_entry.get()
    
    # Data Type Conversion (Fixes "Green Triangle" in Excel)
    try:
        volts = float(volts_entry.get())
    except ValueError:
        volts = volts_entry.get() 

    try:
        freq = float(freq_entry.get())
    except ValueError:
        freq = freq_entry.get()

    full_trace = log_window.get("1.0", tk.END).strip()

    # 2. Check/Create Excel File
    if not os.path.exists(EXCEL_FILENAME):
        wb = Workbook()
        ws = wb.active
        ws.title = "Test History"
        ws.append(["Timestamp", "Technician", "IP Address", "Voltage (V)", "Freq (Hz)", "Log Trace"])
    else:
        try:
            wb = load_workbook(EXCEL_FILENAME)
            ws = wb.active
        except Exception as e:
            messagebox.showerror("File Error", f"Could not load Excel file: {e}")
            return

    # 3. Append Row
    ws.append([timestamp, tech, ip, volts, freq, full_trace])

    # 4. Save
    try:
        wb.save(EXCEL_FILENAME)
        log_message(f"Data appended to {EXCEL_FILENAME}")
        messagebox.showinfo("Success", f"Session saved to {EXCEL_FILENAME}")
    except PermissionError:
        messagebox.showerror("Save Failed", 
            f"Could not save to {EXCEL_FILENAME}!\n\n"
            "Is the Excel file currently open? Please close it and try again."
        )

# --- Main Interface Construction ---
root = tk.Tk()
root.title("Panasonic Critical Facilities | Test Bench")
root.geometry("600x700")

# Header
header_frame = tk.Frame(root, pady=15)
header_frame.pack()
tk.Label(header_frame, text="DC/DC Converter Diagnostic Tool", font=("Segoe UI", 14, "bold")).pack()

# Session Metadata
session_frame = tk.Frame(root, pady=5)
session_frame.pack()
tk.Label(session_frame, text="Technician:", font=("Segoe UI", 10)).pack(side=tk.LEFT, padx=5)
tech_entry = tk.Entry(session_frame, width=25)
tech_entry.pack(side=tk.LEFT, padx=5)

# Connectivity Controls
conn_frame = tk.LabelFrame(root, text="Network Configuration", padx=15, pady=15)
conn_frame.pack(fill="x", padx=15, pady=10)

tk.Label(conn_frame, text="Instrument IP:").grid(row=0, column=0, padx=5)
ip_entry = tk.Entry(conn_frame, width=18)
ip_entry.grid(row=0, column=1, padx=5)

connect_btn = tk.Button(conn_frame, text="Initialize", command=connect_to_scope, bg="#E0E0E0", width=12)
connect_btn.grid(row=0, column=2, padx=10)

status_indicator = tk.Frame(conn_frame, width=16, height=16, bg="gray")
status_indicator.grid(row=0, column=3, padx=5)
status_label = tk.Label(conn_frame, text="Offline", fg="gray", font=("Segoe UI", 9, "bold"))
status_label.grid(row=0, column=4)

sim_mode_var = tk.BooleanVar()
sim_chk = tk.Checkbutton(conn_frame, text="Simulation Mode", var=sim_mode_var, fg="blue")
sim_chk.grid(row=1, column=0, columnspan=2, sticky="w", pady=5)

# Test Parameters
param_frame = tk.LabelFrame(root, text="Signal Parameters", padx=15, pady=15)
param_frame.pack(fill="x", padx=15, pady=5)

tk.Label(param_frame, text="Amplitude (V):").grid(row=0, column=0, padx=5)
volts_entry = tk.Entry(param_frame, width=10)
volts_entry.insert(0, "5.0")
volts_entry.grid(row=0, column=1, padx=5)

tk.Label(param_frame, text="Frequency (Hz):").grid(row=0, column=2, padx=5)
freq_entry = tk.Entry(param_frame, width=10)
freq_entry.insert(0, "50000")
freq_entry.grid(row=0, column=3, padx=5)

# Execution Controls
action_frame = tk.Frame(root, pady=15)
action_frame.pack()

run_btn = tk.Button(action_frame, text="EXECUTE TEST", command=run_test_sequence, 
                    font=("Segoe UI", 10, "bold"), bg="#4CAF50", fg="white", state="disabled", width=18)
run_btn.pack(side=tk.LEFT, padx=10)

stop_btn = tk.Button(action_frame, text="EMERGENCY STOP", command=disable_output, 
                    font=("Segoe UI", 10, "bold"), bg="#F44336", fg="white", state="disabled", width=18)
stop_btn.pack(side=tk.LEFT, padx=10)

# System Log
log_frame = tk.LabelFrame(root, text="Operations Log", padx=5, pady=5)
log_frame.pack(fill="both", expand=True, padx=15, pady=10)
log_window = scrolledtext.ScrolledText(log_frame, height=8, state='disabled', font=("Consolas", 9))
log_window.pack(fill="both", expand=True)

# Footer
footer_frame = tk.Frame(root, pady=10)
footer_frame.pack()
save_btn = tk.Button(footer_frame, text="Update Master Excel Log", command=update_excel_log, width=25, state="disabled")
save_btn.pack()

if __name__ == "__main__":
    load_config()
    root.mainloop()